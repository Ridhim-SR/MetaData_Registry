"""Turn the Kanya Sumangla Portal workbook into Field Dictionary CSVs.

The workbook is one worksheet per scheme, each already shaped like a Field
Dictionary (Dataset Name / Dataset Field / Data Description / Format /
Mandatory / Personal Data, plus Hindi-English labels, availability and a
Sample Value column). This module writes each sheet out as a plain
multi-table Field Dictionary CSV -- the shape parse_field_dictionary and
run_field_dictionary already consume -- dropping everything the pipeline
doesn't read.

Sheet -> department/dataset ownership comes from the workbook's own "Note"
sheet (Scheme / Deptt. / Identifiers): that is the department's own
statement of who owns what, so it is recorded here next to the converter
rather than being re-decided at ingest time. SHEETS is the single source of
truth for B3's ingest loop; the Note sheet is read back in tests to prove
the roster still covers it.

Columns deliberately dropped from the output:
  * Field Label Hindi / Field Label English -- display labels, not metadata.
  * Availability (Public/Internal) -- dataset-level, belongs to a report,
    not to a column definition.
  * Sample Value -- sample production data (Aadhaar/mobile numbers), which
    must not land in a repository or a registry snapshot.
  * Notes / Validation Rules -- free text the parser has nowhere to put.
"""

from __future__ import annotations

import argparse
import csv
from pathlib import Path
from typing import NamedTuple

from src.utils.logger import get_logger

logger = get_logger(__name__)

# Exactly what the parser reads; anything else in the sheet is dropped.
OUTPUT_HEADERS = ["Dataset Name", "Dataset Field", "Data Description", "Format", "Mandatory (Y/N)", "Personal Data (Y/N)"]


class SheetConfig(NamedTuple):
    sheet: str
    scheme: str
    department_id: str
    department_name: str
    dataset: str


# One entry per data sheet in the workbook ("Note" is the roster, not data).
# scheme = the value the Note sheet uses for the department's own sheet.
SHEETS: tuple[SheetConfig, ...] = (
    SheetConfig(
        sheet="Income Cert App Form E District",
        scheme="Income Certificate",
        department_id="revenue_department",
        department_name="Revenue Department (E-District)",
        dataset="income_certificate_app_form_e_district",
    ),
    SheetConfig(
        sheet="Kanya Sumangla Portal",
        scheme="Kanya Sumangla Portal",
        department_id="women_and_child_development_department",
        department_name="Department of Women and Child Development",
        dataset="kanya_sumangla_portal",
    ),
    SheetConfig(
        sheet="PM Matru Vandana Yojana",
        scheme="PM Matru Vandana Yojana",
        department_id="women_and_child_development_department",
        department_name="Department of Women and Child Development",
        dataset="pm_matru_vandana_yojana",
    ),
    SheetConfig(
        sheet="U_win(immunisation)",
        scheme="Immunisation",
        department_id="health_department",
        department_name="Health Department",
        dataset="uwin_immunisation",
    ),
    SheetConfig(
        sheet="UDISE+",
        scheme="UDISE+",
        department_id="basic_education_department",
        department_name="Department of Basic Education",
        dataset="udise_plus",
    ),
    SheetConfig(
        sheet="CM Samuhik Vivah Yojana",
        scheme="Samuhik Vivah",
        department_id="social_welfare_department",
        department_name="Social Welfare Department",
        dataset="cm_samuhik_vivah_yojana",
    ),
)

# The sheet the gitignored samples/kanya_sumangla_field_dictionary.csv is
# built from -- test_field_dictionary_parser's real-sample test pins it to
# these seven cmsvy_* tables.
CMSVY_SHEET = "CM Samuhik Vivah Yojana"


def _load_workbook(path: str | Path):
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - environment guard
        raise RuntimeError("openpyxl is required to read the Kanya Sumangla workbook") from exc
    return openpyxl.load_workbook(str(path), read_only=True, data_only=True)


def sheet_names(workbook_path: str | Path) -> list[str]:
    """Sheet names in workbook order -- used to explain a roster mismatch."""

    workbook = _load_workbook(workbook_path)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def _cell(row: tuple, index: int | None) -> str:
    if index is None or index >= len(row) or row[index] is None:
        return ""
    return str(row[index]).strip()


def _header_index(header: tuple) -> dict[str, int]:
    return {str(name).strip(): i for i, name in enumerate(header) if name is not None and str(name).strip()}


def _sheet_records(sheet) -> tuple[list[dict], int]:
    """Read one sheet's rows into OUTPUT_HEADERS-shaped dicts.

    Columns are looked up by header name, not position: the sheets disagree
    on both extra columns ("Notes" vs "Notes / Validation Rules") and
    trailing empty cells. Returns (records, skipped_rows) where a row is
    skipped when it has no Dataset Name or no Dataset Field -- the workbook
    ships one such name-only row, and the parser would drop it anyway.
    """

    rows = sheet.iter_rows(values_only=True)
    try:
        header = next(rows)
    except StopIteration:
        raise ValueError(f"Sheet '{sheet.title}' is empty -- no header row")

    index = _header_index(header)
    missing = [name for name in OUTPUT_HEADERS if name not in index]
    if missing:
        raise ValueError(f"Sheet '{sheet.title}' is missing column(s) {missing}. Headers found: {list(index)}")

    records: list[dict] = []
    skipped = 0
    for row in rows:
        record = {name: _cell(row, index[name]) for name in OUTPUT_HEADERS}
        if not record["Dataset Name"] or not record["Dataset Field"]:
            skipped += 1
            continue
        records.append(record)
    return records, skipped


def read_sheet(workbook_path: str | Path, sheet_name: str) -> list[dict]:
    """One sheet's field rows as OUTPUT_HEADERS dicts (blank rows dropped)."""

    workbook = _load_workbook(workbook_path)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"No sheet '{sheet_name}' in workbook. Sheets: {workbook.sheetnames}")
        records, skipped = _sheet_records(workbook[sheet_name])
    finally:
        workbook.close()
    if skipped:
        logger.warning(f"kanya_xlsx: '{sheet_name}' skipped {skipped} row(s) with no dataset/field name")
    return records


def read_note_sheet(workbook_path: str | Path) -> dict[str, dict[str, str]]:
    """The workbook's own ownership roster: scheme -> department + identifiers."""

    workbook = _load_workbook(workbook_path)
    try:
        if "Note" not in workbook.sheetnames:
            raise ValueError(f"No 'Note' sheet in workbook. Sheets: {workbook.sheetnames}")
        rows = list(workbook["Note"].iter_rows(values_only=True))
    finally:
        workbook.close()

    header_at = next(
        (i for i, row in enumerate(rows) if any(str(c).strip().lower() == "scheme/schema" for c in row if c)),
        None,
    )
    if header_at is None:
        raise ValueError("Note sheet has no 'Scheme/Schema' header row")

    header = _header_index(rows[header_at])
    roster: dict[str, dict[str, str]] = {}
    for row in rows[header_at + 1 :]:
        scheme = _cell(row, header.get("Scheme/Schema"))
        if not scheme:
            continue
        roster[scheme] = {
            "department": _cell(row, header.get("Deptt.")),
            "identifiers": _cell(row, header.get("Identifiers")),
        }
    return roster


def _write_dictionary(path: Path, records: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=OUTPUT_HEADERS)
        writer.writeheader()
        writer.writerows(records)


def convert_workbook(workbook_path: str | Path, out_dir: str | Path) -> dict[str, Path]:
    """Write one Field Dictionary CSV per data sheet; returns dataset slug -> path."""

    out_dir = Path(out_dir)
    workbook = _load_workbook(workbook_path)
    written: dict[str, Path] = {}
    try:
        for config in SHEETS:
            if config.sheet not in workbook.sheetnames:
                raise ValueError(
                    f"No sheet '{config.sheet}' in workbook. Sheets: {workbook.sheetnames}"
                )
            records, skipped = _sheet_records(workbook[config.sheet])
            path = out_dir / f"{config.dataset}.csv"
            _write_dictionary(path, records)
            written[config.dataset] = path
            logger.info(
                f"kanya_xlsx: '{config.sheet}' -> {path.name} "
                f"({len(records)} field(s), {skipped} skipped row(s))"
            )
    finally:
        workbook.close()
    return written


def write_cmsvy_sample(workbook_path: str | Path, out_path: str | Path) -> Path:
    """Write samples/kanya_sumangla_field_dictionary.csv from the CMSVY sheet."""

    records = read_sheet(workbook_path, CMSVY_SHEET)
    out_path = Path(out_path)
    _write_dictionary(out_path, records)
    logger.info(f"kanya_xlsx: '{CMSVY_SHEET}' -> {out_path} ({len(records)} field(s))")
    return out_path


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, help="path to the Kanya Sumangla Portal workbook")
    parser.add_argument("--out-dir", required=True, help="directory for one Field Dictionary CSV per sheet")
    parser.add_argument(
        "--sample",
        help="also write the CMSVY sheet to this path (samples/kanya_sumangla_field_dictionary.csv)",
    )
    args = parser.parse_args(argv)

    written = convert_workbook(args.input, args.out_dir)
    for dataset, path in written.items():
        print(f"{dataset} -> {path}")
    if args.sample:
        print(f"cmsvy sample -> {write_cmsvy_sample(args.input, args.sample)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
