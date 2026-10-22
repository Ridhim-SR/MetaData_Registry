"""Turn the MLA LADS financial-status workbook into the registry's column CSV.

The portal exports this table as one sheet shaped like::

    row 1   merged title (English in one download, Hindi in the other)
    row 2   "Exported On <timestamp>"
    row 3   the twelve column headers
    row 4+  one row per MLA

Both downloads carry the *same English headers* -- the Hindi workbook differs
only in the title and in the MLA-name / district values -- so the Hindi column
labels are not in the file at all. They come from the portal's own Hindi
rendering and are written into the column CSV's own ``name_Hindi`` column, next
to the English ``name``; OpenMetadata shows them as each column's display name.
They are deliberately *not* appended to the descriptions.

The registry stores column definitions, not rows, so this module maps the
export's prose headers onto stable snake_case names and writes the canonical
``name,name_Hindi,data_type`` CSV that the shared csv parser reads.
"""

from __future__ import annotations

import argparse
import csv
from pathlib import Path

from src.utils.logger import get_logger

logger = get_logger(__name__)

# (header as the export writes it, column name, data type, Hindi name)
#
# The headers are prose ("Name and constituency of the Honorable MLA") and are
# not usable as OpenMetadata column names, so the mapping is fixed here, in one
# reviewable place, instead of being derived from the file on every run.
COLUMNS: tuple[tuple[str, str, str, str], ...] = (
    ("Serial number", "serial_number", "integer", "क्रमसं०"),
    (
        "Name and constituency of the Honorable MLA",
        "mla_name_and_constituency",
        "character varying",
        "माननीय विधायक का नाम व क्षेत्र",
    ),
    ("financial year", "financial_year", "character varying", "वित्तीय वर्ष"),
    ("Nodal District", "nodal_district", "character varying", "नोडल जनपद"),
    (
        "Balance amount of last financial year (in lakhs)",
        "balance_last_financial_year_lakhs",
        "numeric",
        "गत वित्तीय वर्ष की अवशेष धनराशि (लाख मे)",
    ),
    (
        "Amount allocated (in lakh)",
        "amount_allocated_lakhs",
        "numeric",
        "आवंटित धनराशि (लाख मे)",
    ),
    (
        "Total (Remainder + Allocated)   (in lakh)",
        "total_remainder_allocated_lakhs",
        "numeric",
        "योग (अवशेष और आवंटित) (लाख मे)",
    ),
    (
        "Sanctioned Amount (in lakh)",
        "sanctioned_amount_lakhs",
        "numeric",
        "स्वीकृत धनराशि (लाख मे)",
    ),
    (
        "Balance amount for approval (in lakhs)",
        "balance_for_approval_lakhs",
        "numeric",
        "स्वीकृति हेतु अवशेष धनराशि (लाख मे)",
    ),
    (
        "Expenditure Amount (in lakhs)",
        "expenditure_amount_lakhs",
        "numeric",
        "व्यय धनराशि (लाख मे)",
    ),
    (
        "Balance Amount (in lakhs)",
        "balance_amount_lakhs",
        "numeric",
        "अवशेष धनराशि (लाख मे)",
    ),
    (
        "Amount Committed (in lakhs)",
        "amount_committed_lakhs",
        "numeric",
        "प्रतिबद्ध धनराशि (लाख मे)",
    ),
)

# Today's export uses rows 1-2 as a preamble. Searching a few more rows means a
# portal change that adds a line does not need a code change, while a file that
# is not this report still fails loudly instead of being read from row 1.
HEADER_SEARCH_ROWS = 12

_NUMERIC_TYPES = {"integer", "numeric"}


class _NotThisSheet(ValueError):
    """The sheet does not carry the MLA LADS header row."""


def _normalise(value: object) -> str:
    """Headers differ only in case and runs of spaces between exports."""

    return " ".join(str(value).split()).lower() if value is not None else ""


_EXPECTED = {
    _normalise(header): (name, data_type, name_hindi)
    for header, name, data_type, name_hindi in COLUMNS
}
_TYPE_BY_NAME = {name: data_type for _, name, data_type, _ in COLUMNS}
_HINDI_BY_NAME = {name: name_hindi for _, name, _, name_hindi in COLUMNS}


def _load_workbook(path: str | Path):
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - environment guard
        raise RuntimeError("openpyxl is required to read the MLA LADS workbook") from exc

    return openpyxl.load_workbook(str(path), read_only=True, data_only=True)


def _find_header(sheet) -> tuple[int, list[tuple[str, str] | None]]:
    """Return (row number, column per cell) for the sheet's header row.

    Raises when the twelve expected headers are not there -- a silent reshape is
    exactly what this guard exists to stop.
    """

    for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        keys = [_normalise(cell) for cell in row]
        filled = [index for index, key in enumerate(keys) if key]
        known = [index for index in filled if keys[index] in _EXPECTED]
        if len(known) < len(COLUMNS):
            if row_number >= HEADER_SEARCH_ROWS:
                break
            continue

        unknown = [keys[index] for index in filled if keys[index] not in _EXPECTED]
        if unknown:
            raise ValueError(
                f"row {row_number} of '{sheet.title}' has unrecognised column(s) {unknown} -- "
                "the export layout changed, update mla_lads_xlsx.COLUMNS"
            )

        return row_number, [_EXPECTED[keys[i]] if keys[i] else None for i in range(len(keys))]

    raise _NotThisSheet(
        f"no MLA LADS header row in the first {HEADER_SEARCH_ROWS} rows of '{sheet.title}' -- "
        f"expected {len(COLUMNS)} headers starting with {COLUMNS[0][0]!r}"
    )


def read_rows(path: str | Path) -> tuple[list[str], list[list[object]]]:
    """The table as (canonical column names, data rows), preamble skipped."""

    workbook = _load_workbook(path)
    reasons: list[str] = []
    try:
        for sheet in workbook.worksheets:
            try:
                header_row, columns = _find_header(sheet)
            except _NotThisSheet as exc:
                reasons.append(str(exc))
                continue

            indexes = [index for index, column in enumerate(columns) if column is not None]
            names = [columns[index][0] for index in indexes]
            rows = []
            for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
                values = [row[index] if index < len(row) else None for index in indexes]
                if all(value is None or str(value).strip() == "" for value in values):
                    continue
                rows.append(values)

            if not rows:
                raise ValueError(f"'{path}' has the header row but no data rows")
            return names, rows
    finally:
        workbook.close()

    raise ValueError(
        f"'{path}' is not an MLA LADS financial-status export: " + " | ".join(reasons)
    )


def _check_values(names: list[str], rows: list[list[object]]) -> None:
    """Numeric columns must actually hold numbers -- catches a shifted column."""

    for index, name in enumerate(names):
        if _TYPE_BY_NAME.get(name) not in _NUMERIC_TYPES:
            continue
        for offset, row in enumerate(rows, start=1):
            value = row[index]
            if value is None or str(value).strip() == "":
                continue
            try:
                float(value)
            except (TypeError, ValueError):
                raise ValueError(
                    f"column '{name}' row {offset} holds {value!r}, which is not a number"
                ) from None


def schema_columns(path: str | Path) -> list[dict]:
    """Column definitions in the shape the shared csv parser produces."""

    names, rows = read_rows(path)
    _check_values(names, rows)
    return [
        {
            "name": name,
            "name_hindi": _HINDI_BY_NAME[name],
            "data_type": _TYPE_BY_NAME[name],
            "length": None,
            "scale": None,
            "nullable": True,
            "default": None,
        }
        for name in names
    ]


def write_schema_csv(workbook_path: str | Path, out_path: str | Path) -> list[dict]:
    """Write the canonical ``name,name_Hindi,data_type`` CSV the parser reads."""

    columns = schema_columns(workbook_path)
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with out_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["name", "name_Hindi", "data_type"])
        for column in columns:
            writer.writerow([column["name"], column["name_hindi"], column["data_type"]])
    logger.info(f"mla_lads_xlsx: {len(columns)} column(s) -> {out_path}")
    return columns


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--input", required=True, help="path to the MLA LADS workbook")
    parser.add_argument("--out", required=True, help="path for the column-definition CSV")
    args = parser.parse_args(argv)

    columns = write_schema_csv(args.input, args.out)
    for column in columns:
        print(f"{column['name']:<34} {column['name_hindi']:<44} {column['data_type']}")
    print(f"\n{len(columns)} column(s) -> {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
