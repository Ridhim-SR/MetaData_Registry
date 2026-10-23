"""Turn an arbitrary department workbook into the registry's column CSV.

The ingest API advertises that it takes "the department's raw file", and most
departments send .xlsx. The registry itself stores column *definitions*, so a
workbook is reduced to ``name[,name_Hindi],data_type`` here.

Unlike the dataset-specific converters (``mla_lads_xlsx``, ``kanya_xlsx``),
this one does not know the department's prose headers. It therefore:

* skips preamble rows (a merged title, "Exported On ...") and picks the row
  with the most filled cells in the first ``HEADER_SEARCH_ROWS`` rows;
* turns ASCII headers into snake_case names, and keeps a non-ASCII header as
  the column's ``name_Hindi`` (the column name becomes ``column_<n>``);
* infers a type from the values below the header (integer / double precision /
  boolean / date / timestamp / character varying).

Dataset-specific converters stay the better choice where they exist: the API
uses ``mla_lads_xlsx`` for ``mla_lads`` so its reviewed Hindi labels survive.
"""

from __future__ import annotations

import csv
import datetime as dt
import re
from decimal import Decimal
from pathlib import Path

from src.utils.logger import get_logger

logger = get_logger(__name__)

HEADER_SEARCH_ROWS = 20
TYPE_SAMPLE_ROWS = 200


def _load_workbook(path: str | Path):
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - environment guard
        raise RuntimeError("openpyxl is required to read an Excel workbook") from exc
    return openpyxl.load_workbook(str(path), read_only=True, data_only=True)


def _filled(row: tuple) -> int:
    return sum(1 for value in row if value is not None and str(value).strip() != "")


def _find_header(sheet) -> tuple[int, list[object]]:
    """(1-based row number, that row's values) for the sheet's header row."""

    best: tuple[int, list[object]] | None = None
    best_filled = 0
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if row_number > HEADER_SEARCH_ROWS:
            break
        filled = _filled(row)
        # A merged title or a one-cell "Exported On ..." row must not win; the
        # header is the fullest row in the preamble. Ties go to the earliest,
        # which is the header when data rows also fill every column.
        if filled >= 2 and filled > best_filled:
            best, best_filled = (row_number, list(row)), filled
    if best is None:
        raise ValueError(
            f"'{sheet.title}' has no header row in its first {HEADER_SEARCH_ROWS} rows "
            "(need one row with at least two column headings)"
        )
    return best


def _column_name(header: object, position: int, used: set[str]) -> tuple[str, str]:
    """(column name, Hindi/original label). The label is kept only when the
    header has no usable ASCII spelling."""

    text = "" if header is None else str(header).strip()
    slug = re.sub(r"[^A-Za-z0-9]+", "_", text).strip("_").lower()
    if slug and not slug[0].isdigit():
        name, label = slug, ""
    elif slug:
        name, label = f"col_{slug}", ""
    else:
        name, label = f"column_{position}", text
    candidate, suffix = name, 2
    while candidate in used:
        candidate = f"{name}_{suffix}"
        suffix += 1
    used.add(candidate)
    return candidate, label


def _is_number(value: object) -> bool:
    return isinstance(value, (int, float, Decimal)) and not isinstance(value, bool)


def _infer_type(values: list[object]) -> str:
    samples = [value for value in values if value is not None and str(value).strip() != ""]
    if not samples:
        return "character varying"
    if all(isinstance(value, bool) for value in samples):
        return "boolean"
    if all(isinstance(value, int) and not isinstance(value, bool) for value in samples):
        return "integer"
    if all(_is_number(value) for value in samples):
        return "double precision"
    if all(isinstance(value, dt.datetime) for value in samples):
        return "timestamp"
    if all(isinstance(value, dt.date) for value in samples):
        return "date"
    return "character varying"


def _choose_sheet(workbook, wanted: str | None):
    if wanted:
        if wanted not in workbook.sheetnames:
            raise ValueError(f"worksheet '{wanted}' not found -- this workbook has: {', '.join(workbook.sheetnames)}")
        return workbook[wanted]
    best = None
    for sheet in workbook.worksheets:
        try:
            header_row, header = _find_header(sheet)
        except ValueError:
            continue
        width = len([cell for cell in header if cell is not None and str(cell).strip() != ""])
        if best is None or width > best[1]:
            best = (sheet, width, header_row, header)
    if best is None:
        raise ValueError(
            "no worksheet has a usable header row "
            f"(sheets: {', '.join(workbook.sheetnames) or 'none'})"
        )
    return best[0]


def write_schema_csv(
    workbook_path: str | Path,
    out_path: str | Path,
    sheet: str | None = None,
) -> dict:
    """Write the canonical column CSV and return a small report:
    ``{"sheet": ..., "header_row": ..., "columns": n, "rows_sampled": n}``."""

    workbook_path, out_path = Path(workbook_path), Path(out_path)
    workbook = _load_workbook(workbook_path)
    try:
        chosen = _choose_sheet(workbook, sheet)
        header_row, header = _find_header(chosen)

        positions = [index for index, cell in enumerate(header) if cell is not None and str(cell).strip() != ""]
        used: set[str] = set()
        columns: list[tuple[str, str]] = [
            _column_name(header[index], position + 1, used) for position, index in enumerate(positions)
        ]
        if not columns:
            raise ValueError(f"'{chosen.title}' has no column headings in row {header_row}")

        samples: list[list[object]] = []
        for row in chosen.iter_rows(min_row=header_row + 1, values_only=True):
            if _filled(row) == 0:
                continue
            samples.append([row[index] if index < len(row) else None for index in positions])
            if len(samples) >= TYPE_SAMPLE_ROWS:
                break
        if not samples:
            raise ValueError(f"'{chosen.title}' has the header row but no data rows below it")

        types = [_infer_type([row[index] for row in samples]) for index in range(len(columns))]

        out_path.parent.mkdir(parents=True, exist_ok=True)
        with out_path.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.writer(handle)
            with_hindi = any(label for _, label in columns)
            writer.writerow(["name", "name_Hindi", "data_type"] if with_hindi else ["name", "data_type"])
            for (name, label), data_type in zip(columns, types):
                writer.writerow([name, label, data_type] if with_hindi else [name, data_type])

        report = {"sheet": chosen.title, "header_row": header_row, "columns": len(columns), "rows_sampled": len(samples)}
        logger.info(
            f"xlsx_schema: '{workbook_path.name}' -> {len(columns)} column(s) "
            f"from sheet '{chosen.title}' (header row {header_row})"
        )
        return report
    finally:
        workbook.close()
